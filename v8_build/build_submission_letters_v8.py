# -*- coding: utf-8 -*-
"""Convert the v8 cover letter and title page markdown to .docx (simple converter)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SUB = ROOT / "01_submission"

TARGETS = [
    ("Applied_Sciences_Cover_Letter_v8.md", "Applied_Sciences_Cover_Letter_v8.docx"),
    ("Applied_Sciences_Title_Page_v8.md", "Applied_Sciences_Title_Page_v8.docx"),
]

SUPER = {"1": "\u00b9", "2": "\u00b2", "3": "\u00b3", "4": "\u2074",
         "5": "\u2075", "6": "\u2076", "7": "\u2077", "8": "\u2078", "9": "\u2079"}


def add_runs(par, text: str) -> None:
    for token in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|\[\d\]|\d\u002a)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            r = par.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            r = par.add_run(token[1:-1])
            r.italic = True
        elif re.fullmatch(r"\[\d\]", token):
            r = par.add_run(SUPER[token[1]])
            r.font.superscript = True
        else:
            par.add_run(token)


def convert(md_name: str, docx_name: str) -> Path:
    lines = (SUB / md_name).read_text(encoding="utf8").splitlines()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(14)
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(line[3:])
            r.bold = True
            r.font.size = Pt(12)
            continue
        if line.startswith("**") and line.rstrip().endswith("**") and line.count("**") == 2:
            p = doc.add_paragraph()
            r = p.add_run(line.strip("*"))
            r.bold = True
            continue
        p = doc.add_paragraph()
        add_runs(p, line)

    out = SUB / docx_name
    doc.save(out)
    print("saved", out, out.stat().st_size)
    return out


if __name__ == "__main__":
    for md, dx in TARGETS:
        convert(md, dx)
